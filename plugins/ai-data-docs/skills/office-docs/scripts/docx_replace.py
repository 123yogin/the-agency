#!/usr/bin/env python3
"""Find-and-replace in a .docx that survives Word splitting text across runs.

Usage:
    python3 docx_replace.py in.docx out.docx "old text" "new text" [--all]

Searches body paragraphs, tables (including nested cells) and headers/footers.
The replacement takes the formatting of the run where the match starts.
Prints the number of replacements; exits 1 if there were none.
Requires: pip install python-docx
"""
import sys
from docx import Document


def replace_in_paragraph(paragraph, old, new):
    runs = paragraph.runs
    full = "".join(r.text for r in runs)
    start = full.find(old)
    if start == -1:
        return False
    end = start + len(old)
    pos, placed = 0, False
    for r in runs:
        r_start, r_end = pos, pos + len(r.text)
        pos = r_end
        if r_end <= start or r_start >= end:
            continue
        before = r.text[: max(0, start - r_start)]
        after = r.text[end - r_start:] if r_end > end else ""
        r.text = before + ("" if placed else new) + after
        placed = True
    return True


def iter_paragraphs(container):
    for p in container.paragraphs:
        yield p
    for table in getattr(container, "tables", []):
        for row in table.rows:
            for cell in row.cells:
                yield from iter_paragraphs(cell)


def main(argv):
    if len(argv) < 5:
        print(__doc__)
        return 2
    src, dst, old, new = argv[1:5]
    replace_all = "--all" in argv[5:]
    doc = Document(src)
    containers = [doc]
    for section in doc.sections:
        containers += [section.header, section.footer]
    count = 0
    for container in containers:
        for p in iter_paragraphs(container):
            while replace_in_paragraph(p, old, new):
                count += 1
                if not replace_all or old in new:
                    break
            if count and not replace_all:
                break
        if count and not replace_all:
            break
    doc.save(dst)
    # verify by reopening
    check = Document(dst)
    remaining = sum(p.text.count(old) for c in [check] for p in iter_paragraphs(c))
    print(f"replacements: {count}; occurrences of old text left in body: {remaining}")
    return 0 if count else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv))
